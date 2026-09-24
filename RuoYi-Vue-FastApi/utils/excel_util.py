"""
Excel导出/导入工具（对应java版ExcelUtil，基于openpyxl）

规范：POST /xxx/export 接口用 export_excel 返回xlsx流；列定义参照java实体的@Excel注解
导入用 parse_excel 按表头标题解析为dict列表
"""
import io
from datetime import datetime, date
from typing import Any, List, Optional
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from fastapi.responses import StreamingResponse
from utils.common_util import _format_value


class ExcelColumn:
    """
    导出列定义（对应java版@Excel注解）
    """

    def __init__(self, title: str, field: str,
                 dict_type: Optional[dict] = None,
                 width: int = 16):
        """
        :param title: 表头标题
        :param field: 数据字段名（驼峰，与前端/transform_result一致）
        :param dict_type: 值转换映射（如 {'0': '男', '1': '女'}，对应java dictType转换）
        :param width: 列宽
        """
        self.title = title
        self.field = field
        self.dict_type = dict_type
        self.width = width


def _cell_value(row: Any, col: ExcelColumn):
    """
    取单元格值：支持dict与对象；dict_type映射转换；日期格式化
    """
    if isinstance(row, dict):
        value = row.get(col.field)
    else:
        value = getattr(row, col.field, None)
    if col.dict_type and value is not None:
        return col.dict_type.get(str(value), value)
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    return value


def build_workbook(sheet_name: str, columns: List[ExcelColumn], rows: List[Any]) -> Workbook:
    """
    构建工作簿：表头样式对齐java版（加粗白字+灰底）
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='808080', end_color='808080', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center')

    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=col.title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = col.width

    for r, row in enumerate(rows, start=2):
        for c, col in enumerate(columns, start=1):
            ws.cell(row=r, column=c, value=_cell_value(row, col))

    return wb


def export_excel(sheet_name: str, columns: List[ExcelColumn], rows: List[Any],
                 filename: str) -> StreamingResponse:
    """
    生成xlsx下载流（对齐java版exportExcel的响应头）
    """
    wb = build_workbook(sheet_name, columns, rows)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    from urllib.parse import quote
    encoded = quote(f'{filename}.xlsx')
    return StreamingResponse(
        buf,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f"attachment; filename*=utf-8''{encoded}",
                 'Access-Control-Expose-Headers': 'Content-Disposition'}
    )


def export_template(sheet_name: str, columns: List[ExcelColumn], filename: str,
                    example_rows: Optional[List[dict]] = None) -> StreamingResponse:
    """
    导入模板下载（表头 + 可选示例行，对应java版importTemplateExcel）
    """
    rows = example_rows or []
    return export_excel(sheet_name, columns, rows, filename)


def parse_excel(file_bytes: bytes, columns: List[ExcelColumn]) -> List[dict]:
    """
    解析上传的xlsx为dict列表（按表头标题映射字段，对应java版importExcel）
    :param file_bytes: 上传文件内容
    :param columns: 列定义（title->field映射）
    :return: [{field: value}, ...]
    """
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    title_to_field = {col.title: col for col in columns}

    # 定位表头行（第一行包含已知标题的行）
    header_row = None
    for row in rows_iter:
        if any(isinstance(v, str) and v in title_to_field for v in row):
            header_row = row
            break
    if header_row is None:
        wb.close()
        raise ValueError('未识别到有效的表头行')

    # 列下标 -> 列定义
    idx_columns = []
    for idx, title in enumerate(header_row):
        if isinstance(title, str) and title in title_to_field:
            idx_columns.append((idx, title_to_field[title]))

    result = []
    for row in rows_iter:
        if all(v is None or str(v).strip() == '' for v in row):
            continue  # 跳过空行
        item = {}
        for idx, col in idx_columns:
            value = row[idx] if idx < len(row) else None
            # dict_type映射：导入时是"翻译后"的值还是编码？java导入按readConverterExp反查
            raw = _format_value(value)
            if col.dict_type and raw is not None:
                # 反查：'男'->'0'；查不到时保留原值（用户可能直接填编码）
                reversed_map = {v: k for k, v in col.dict_type.items()}
                item[col.field] = reversed_map.get(str(raw), str(raw))
            else:
                item[col.field] = raw
        result.append(item)
    wb.close()
    return result
