"""
spec-01 基础设施单元测试（不依赖真实db/redis的部分）
"""
import sys
import os

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from utils.page_util import _safe_order_column, PageDomain
from module_admin.aspect.interface_auth import _match
from module_admin.annotation.log_annotation import _sanitize, _truncate
from utils.xss_util import _clean_obj
from utils.common_util import snake_to_camel, transform_result
from common.enums import BusinessType, UserStatus, OperatorType
from common.message_util import message


class TestOrderColumn:
    """分页排序字段白名单"""

    def test_camel_to_snake(self):
        assert _safe_order_column('userId') == 'user_id'
        assert _safe_order_column('createTime') == 'create_time'
        assert _safe_order_column('user_name') == 'user_name'

    def test_injection_rejected(self):
        assert _safe_order_column('userId; DROP TABLE x') is None
        assert _safe_order_column("user_id' --") is None
        assert _safe_order_column('1) UNION') is None
        assert _safe_order_column('') is None
        assert _safe_order_column(None) is None

    def test_order_clause(self):
        d = PageDomain(1, 10, 'createTime', 'desc')
        assert d.order_clause == 'create_time desc'
        d2 = PageDomain(1, 10, 'hack--', 'asc')
        assert d2.order_clause is None

    def test_offset(self):
        assert PageDomain(3, 10).offset == 20


class TestPermissionMatch:
    """权限字符串匹配（对齐java PermissionService.hasPermi）"""

    def test_exact(self):
        assert _match({'system:user:list'}, 'system:user:list')

    def test_all_permission(self):
        assert _match({'*:*:*'}, 'monitor:job:list')

    def test_segment_wildcard(self):
        assert _match({'system:user:*'}, 'system:user:list')
        assert _match({'system:*:list'}, 'system:user:list')
        assert not _match({'system:user:*'}, 'monitor:job:list')

    def test_no_match(self):
        assert not _match({'system:user:list'}, 'monitor:job:list')
        assert not _match(set(), 'system:user:list')


class TestLogSanitize:
    """操作日志脱敏"""

    def test_password_masked(self):
        data = {'username': 'a', 'password': 'secret', 'newPassword': 'x', 'oldPassword': 'y'}
        out = _sanitize(data)
        assert out['password'] == '*' and out['newPassword'] == '*'
        assert out['username'] == 'a'

    def test_nested_masked(self):
        out = _sanitize({'items': [{'confirmPassword': 'z'}]})
        assert out['items'][0]['confirmPassword'] == '*'

    def test_truncate(self):
        assert len(_truncate('a' * 3000)) == 2000
        assert _truncate(None) == ''


class TestXssClean:
    """XSS清洗"""

    def test_tag_stripped(self):
        # 标签被剥离，标签内文本保留（对齐java HTMLFilter语义）
        assert _clean_obj({'a': '<script>alert(1)</script>hi'}) == {'a': 'alert(1)hi'}
        assert _clean_obj({'a': 'plain'}) == {'a': 'plain'}
        assert _clean_obj({'a': '<img src=x onerror=alert(1)>'}) == {'a': ''}

    def test_nested(self):
        out = _clean_obj({'body': ['<b>x</b>', {'c': '<i>y</i>'}]})
        assert out['body'][0] == 'x' and out['body'][1]['c'] == 'y'


class TestCamelCase:
    """驼峰转换"""

    def test_basic(self):
        assert snake_to_camel('user_name') == 'userName'
        assert snake_to_camel('userName') == 'userName'
        assert snake_to_camel('id') == 'id'

    def test_transform_dict(self):
        out = transform_result({'user_name': 'a'})
        assert out == {'userName': 'a'}


class TestEnumsAgainstJava:
    """枚举值与java版对表"""

    def test_business_type_codes(self):
        # java @Log business_type 落库数字
        assert BusinessType.OTHER.code == 0
        assert BusinessType.INSERT.code == 1
        assert BusinessType.UPDATE.code == 2
        assert BusinessType.DELETE.code == 3
        assert BusinessType.GRANT.code == 4
        assert BusinessType.EXPORT.code == 5
        assert BusinessType.IMPORT.code == 6
        assert BusinessType.FORCE.code == 7
        assert BusinessType.GENCODE.code == 8
        assert BusinessType.CLEAN.code == 9

    def test_user_status_codes(self):
        assert UserStatus.OK.code == '0'
        assert UserStatus.DISABLE.code == '1'
        assert UserStatus.DELETED.code == '2'

    def test_operator_type_codes(self):
        assert OperatorType.OTHER.code == 0
        assert OperatorType.MANAGE.code == 1
        assert OperatorType.MOBILE.code == 2


class TestMessage:
    """文案占位符"""

    def test_placeholder(self):
        assert message('user.password.retry.limit.exceed', 5, 10) == '密码输入错误5次，帐户锁定10分钟'

    def test_missing_key(self):
        assert message('no.such.key') == 'no.such.key'


class TestExcel:
    """Excel导出"""

    def test_build_workbook(self):
        from openpyxl import load_workbook
        import io
        from utils.excel_util import ExcelColumn, build_workbook
        cols = [ExcelColumn('用户编号', 'userId'), ExcelColumn('用户名称', 'userName')]
        rows = [{'userId': 1, 'userName': 'admin'}, {'userId': 2, 'userName': 'ry'}]
        wb = build_workbook('测试', cols, rows)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        wb2 = load_workbook(buf)
        ws = wb2.active
        assert ws.cell(1, 1).value == '用户编号'
        assert ws.cell(2, 2).value == 'admin'
        assert ws.max_row == 3

    def test_dict_mapping(self):
        from utils.excel_util import ExcelColumn, _cell_value
        col = ExcelColumn('性别', 'sex', dict_type={'0': '男', '1': '女', '2': '未知'})
        assert _cell_value({'sex': '0'}, col) == '男'
        assert _cell_value({'sex': '9'}, col) == '9'


if __name__ == '__main__':
    sys.exit(pytest.main([__file__, '-v']))


class TestInfraSpecLeftovers:
    """spec-01 遗留项补充测试：403权限、data_scope条件、Lua限流脚本"""

    def test_permission_403_logic(self):
        """common角色权限不含system:user:list时被_match拒绝；admin通配放行"""
        from module_admin.aspect.interface_auth import _match
        common_perms = {'system:user:query', 'system:user:add'}
        assert not _match(common_perms, 'system:user:list')
        assert _match({'*:*:*'}, 'system:user:list')

    def test_data_scope_admin_no_filter(self):
        """admin用户数据权限为空条件（不过滤）"""
        import asyncio
        from module_admin.aspect.data_scope import build_data_scope_conditions, ADMIN_USER_ID
        from module_admin.entity.do.entity import SysDept, SysUser

        async def run():
            conds = await build_data_scope_conditions(
                {'user_id': ADMIN_USER_ID}, SysDept, SysUser)
            return conds == []

        assert asyncio.run(run())

    def test_rate_limit_lua(self):
        """限流Lua脚本在fakeredis下的计数与TTL逻辑"""
        import asyncio
        from module_admin.annotation.repeat_rate_limit import RATE_LIMIT_LUA

        async def run():
            import fakeredis.aioredis
            r = fakeredis.aioredis.FakeRedis(decode_responses=True)
            key = 'rate_limit:test'
            vals = []
            for _ in range(3):
                v = await r.eval(RATE_LIMIT_LUA, 1, key, 3, 60)
                vals.append(int(v))
            v4 = int(await r.eval(RATE_LIMIT_LUA, 1, key, 3, 60))
            ttl = await r.ttl(key)
            await r.aclose()
            assert vals == [1, 2, 3], vals
            assert v4 > 3, v4
            assert ttl > 0
            return True

        assert asyncio.run(run())
